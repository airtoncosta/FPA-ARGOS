import openpyxl

excel_file = r"C:\Users\Airton\OneDrive\Área de Trabalho\CNES_Profissionais_SÃO_LUÍS_GONZAGA_DO_MARANHÃO_202608.xlsx"
wb = openpyxl.load_workbook(excel_file)
print("Planilhas no arquivo:", wb.sheetnames)

ws1 = wb["Colaboradores & Vínculos"]
print("Total de linhas na aba 'Colaboradores & Vínculos':", ws1.max_row - 1)

caps_rows = []
for r in range(2, ws1.max_row + 1):
    cnes = str(ws1.cell(row=r, column=1).value or "").strip()
    if cnes == "9602550":
        caps_rows.append({
            "nome": ws1.cell(row=r, column=3).value,
            "cns": ws1.cell(row=r, column=4).value,
            "cbo": ws1.cell(row=r, column=5).value,
            "ocup": ws1.cell(row=r, column=6).value,
            "ch": ws1.cell(row=r, column=10).value,
            "situacao": ws1.cell(row=r, column=14).value
        })

print(f"\nTotal de profissionais no CAPS (CNES 9602550): {len(caps_rows)}")
for i, p in enumerate(caps_rows, 1):
    print(f"  {i:02d}. {p['nome']} | CNS: {p['cns']} | {p['ocup']} | {p['ch']}h | {p['situacao']}")

ws2 = wb["Estabelecimentos"]
print(f"\nTotal de estabelecimentos na aba 'Estabelecimentos': {ws2.max_row - 1}")
for r in range(2, ws2.max_row + 1):
    cnes = str(ws2.cell(row=r, column=1).value or "").strip()
    nome = ws2.cell(row=r, column=2).value
    qtd = ws2.cell(row=r, column=8).value
    if cnes == "9602550":
        print(f"  -> Verificação CAPS na lista de Unidades: CNES {cnes} | {nome} | Qtd: {qtd}")

ws3 = wb["Auditoria Portaria 134"]
print(f"\nTotal de apontamentos na aba 'Auditoria Portaria 134': {ws3.max_row - 1}")
